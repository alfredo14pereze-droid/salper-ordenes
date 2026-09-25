#!/usr/bin/env python3
"""
V66 — Producción, Fase 1: importa los CATÁLOGOS del Excel a SQL.

Lee scripts/data/Salper_Produccion.xlsm (ignorado por git: trae nombres y datos
de pago del personal) y GENERA scripts/data/import_fase1.sql, que se pega en el
SQL Editor de Supabase. No se conecta a la base ni necesita llaves.

Carga: configuración, 635 operaciones, 35 operadoras (con su número de
operadora tomado de CALCULO PRODUCCION) y las reglas de premios (hoja TABLAS).
NO importa la captura semanal ni el historial (eso es la Fase 2).

IDEMPOTENTE: todo va con `on conflict do nothing`; correr el SQL dos veces no
duplica nada ni pisa ediciones hechas después desde la app.

Limpieza: se recortan espacios de todos los textos y se corrige el nombre de
prenda "CHAMARA CON FORRO" -> "CHAMARRA CON FORRO".

Uso:  python3 scripts/import_produccion_fase1.py
Requiere: pip install openpyxl
"""
import datetime as dt
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSM = os.path.join(ROOT, "scripts", "data", "Salper_Produccion.xlsm")
OUT = os.path.join(ROOT, "scripts", "data", "import_fase1.sql")

PRENDA_FIXES = {"CHAMARA CON FORRO": "CHAMARRA CON FORRO"}


def q(v):
    """Literal SQL de texto (o null)."""
    if v is None:
        return "null"
    return "'" + str(v).replace("'", "''") + "'"


def clean(v):
    return None if v is None else " ".join(str(v).split())


def num(v):
    return repr(float(v)) if isinstance(v, float) and not float(v).is_integer() else str(int(v))


def main():
    if not os.path.exists(XLSM):
        sys.exit(f"No encuentro {XLSM}")
    wb = openpyxl.load_workbook(XLSM, data_only=True)
    lines = [
        "-- Generado por scripts/import_produccion_fase1.py el " + dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "-- Idempotente (on conflict do nothing). NO subir a git: trae nombres del personal.",
        "begin;",
        "",
    ]

    # --- Configuración (encabezado de OPERACIONES: H2 y H3)
    ws = wb["OPERACIONES"]
    precio, segundos_jornada = ws["H2"].value, ws["H3"].value
    assert precio == 0.025 and segundos_jornada == 25650, (precio, segundos_jornada)
    lines += [
        "insert into public.prod_config (clave, valor) values",
        f"  ('precio_por_segundo', {precio}), ('segundos_jornada', {segundos_jornada})",
        "on conflict (clave) do nothing;",
        "",
    ]

    # --- Operaciones (tabla Operaciones: B7:F641)
    ops = []
    for folio, prenda, parte, operacion, seg in ws.iter_rows(min_row=7, max_row=641, min_col=2, max_col=6, values_only=True):
        prenda = clean(prenda)
        prenda = PRENDA_FIXES.get(prenda, prenda)
        ops.append((int(folio), prenda, clean(parte), clean(operacion), seg))
    assert len(ops) == 635 and len({o[0] for o in ops}) == 635, "se esperaban 635 folios únicos"
    lines.append("insert into public.prod_operaciones (folio, prenda, parte, operacion, segundos) values")
    lines.append(",\n".join(f"  ({f}, {q(p)}, {q(pa)}, {q(o)}, {num(s)})" for f, p, pa, o, s in ops))
    lines += ["on conflict (folio) do nothing;", ""]

    # --- Operadoras (EMPLEADOS + número de operadora de CALCULO PRODUCCION)
    numero = {}
    for folio, oper in wb["CALCULO PRODUCCION"].iter_rows(min_row=4, max_row=886, min_col=1, max_col=2, values_only=True):
        if isinstance(folio, str) and folio.startswith("EMP") and oper is not None:
            numero[folio.strip()] = int(oper)
    we = wb["EMPLEADOS"]
    emps = [r for r in we.iter_rows(min_row=2, max_row=we.max_row, min_col=1, max_col=5, values_only=True) if r[0]]
    assert len(emps) == 35, "se esperaban 35 personas"
    rows = []
    for folio, nombre, puesto, participa, activo in emps:
        folio = clean(folio)
        rows.append(
            f"  ({q(folio)}, {numero.get(folio, 'null')}, {q(clean(nombre))}, {q(clean(puesto))}, "
            f"{'true' if clean(participa) == 'SI' else 'false'}, {'true' if clean(activo) == 'SI' else 'false'})"
        )
    lines.append("insert into public.prod_operadoras (folio_empleado, numero_operadora, nombre, puesto, participa_bonos, activo) values")
    lines.append(",\n".join(rows))
    lines += ["on conflict (folio_empleado) do nothing;", ""]

    # --- Reglas de premios (hoja TABLAS; 3 bloques separados por filas vacías)
    reglas, tipo = [], None
    for a, b in wb["TABLAS"].iter_rows(min_row=1, max_row=40, min_col=1, max_col=2, values_only=True):
        if isinstance(a, str):
            tipo = {"Desde": "meta", "Desde Lugar": "lugar", "Desde %": "mejora"}.get(a.strip())
        elif a is not None and tipo:
            reglas.append((tipo, a, b))
    lines.append("insert into public.prod_reglas_premios (tipo, desde, bono) values")
    lines.append(",\n".join(f"  ({q(t)}, {num(d)}, {num(b)})" for t, d, b in reglas))
    lines += ["on conflict (tipo, desde) do nothing;", ""]

    lines += ["commit;", ""]
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    por_tipo = {t: sum(1 for r in reglas if r[0] == t) for t in ("meta", "lugar", "mejora")}
    print(f"OK -> {OUT}")
    print(f"  operaciones: {len(ops)} | operadoras: {len(emps)} (con número: {len(numero)}) | reglas: {len(reglas)} {por_tipo}")
    sin_numero = [clean(e[0]) + " " + clean(e[1]) for e in emps if clean(e[0]) not in numero]
    print(f"  sin número de operadora: {sin_numero}")


if __name__ == "__main__":
    main()
