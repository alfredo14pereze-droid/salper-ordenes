#!/usr/bin/env python3
"""
V67 — Producción, Fase 2: importa el HISTORIAL semanal del Excel a SQL.

Lee la hoja `HISTORIAL PRODUCCION` (tblHistorial) de scripts/data/Salper_Produccion.xlsm
y GENERA scripts/data/import_historial.sql + scripts/data/mapeo_semanas.txt.

Regla de semanas: la semana de producción va de MIÉRCOLES a MARTES. Las fechas
del Excel son irregulares (martes, miércoles, jueves): cada fecha se asigna a la
semana cuyo MARTES de cierre es el más reciente <= a esa fecha. Si dos fechas
caen en la misma semana (el Excel guardó la misma semana dos veces), se conserva
SOLO la más reciente (la corrección) y se descarta la otra.

Cada semana importada queda `importada = true` y `estado = 'aprobada'`. El valor
de cada persona va a prod_valor_semana. Idempotente (`on conflict do nothing`).

Uso:  python3 scripts/import_produccion_historial.py     (requiere openpyxl)
"""
import collections
import datetime as dt
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSM = os.path.join(ROOT, "scripts", "data", "Salper_Produccion.xlsm")
OUT_SQL = os.path.join(ROOT, "scripts", "data", "import_historial.sql")
OUT_MAP = os.path.join(ROOT, "scripts", "data", "mapeo_semanas.txt")
DIAS = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]


def martes_de_cierre(d):
    """Martes más reciente <= d (weekday: lunes=0 ... martes=1)."""
    return d - dt.timedelta(days=(d.weekday() - 1) % 7)


def main():
    wb = openpyxl.load_workbook(XLSM, data_only=True)
    ws = wb["HISTORIAL PRODUCCION"]
    por_fecha = collections.OrderedDict()
    for fecha, folio, _op, _nombre, prod, _sem in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if fecha is None:
            continue
        por_fecha.setdefault(fecha.date(), []).append((str(folio).strip(), round(float(prod or 0), 4)))

    # fecha -> semana (cierre)
    asignada = {d: martes_de_cierre(d) for d in por_fecha}
    grupos = collections.defaultdict(list)
    for d, fin in asignada.items():
        grupos[fin].append(d)
    conservar = {fin: max(fechas) for fin, fechas in grupos.items()}  # la más reciente

    lineas = ["FECHA ORIGINAL | DÍA | SEMANA ASIGNADA (miércoles → martes) | PERSONAS | TOTAL VALOR GENERADO | ACCIÓN"]
    for d in sorted(por_fecha):
        fin = asignada[d]
        ini = fin - dt.timedelta(days=6)
        filas = por_fecha[d]
        total = round(sum(v for _, v in filas), 2)
        accion = "IMPORTAR" if conservar[fin] == d else f"DESCARTAR (duplicada; se conserva la del {conservar[fin]})"
        lineas.append(f"{d} | {DIAS[d.weekday()]} | {ini} → {fin} | {len(filas)} | {total:,.2f} | {accion}")
    faltantes = []
    fins = sorted(conservar)
    for a, b in zip(fins, fins[1:]):
        gap = (b - a).days // 7 - 1
        for k in range(1, gap + 1):
            faltantes.append(a + dt.timedelta(days=7 * k))
    lineas.append("")
    lineas.append(f"Semanas a crear: {len(conservar)}. Semanas sin datos entre medias (cierre martes): {[str(x) for x in faltantes]}")
    with open(OUT_MAP, "w", encoding="utf-8") as f:
        f.write("\n".join(lineas) + "\n")

    sql = [
        "-- Generado por scripts/import_produccion_historial.py. Idempotente. NO subir a git.",
        "begin;",
        "insert into public.prod_semanas (fecha_inicio, fecha_fin, estado, importada, notas) values",
    ]
    filas_sem = []
    for fin in sorted(conservar):
        ini = fin - dt.timedelta(days=6)
        orig = conservar[fin]
        filas_sem.append(f"  ('{ini}', '{fin}', 'aprobada', true, 'Importada del Excel (fecha original {orig})')")
    sql += [",\n".join(filas_sem), "on conflict (fecha_inicio) do nothing;", ""]
    # Formato compacto: una fila por semana con "EMP001:valor,EMP002:valor,..."
    filas = []
    for fin in sorted(conservar):
        datos = ",".join(f"{folio}:{v}" for folio, v in por_fecha[conservar[fin]])
        filas.append(f"  ('{fin}'::date, '{datos}')")
    sql += [
        "insert into public.prod_valor_semana (semana_id, operadora_id, valor_generado)",
        "select s.id, o.id, split_part(kv, ':', 2)::numeric",
        "from (values",
        ",\n".join(filas),
        ") as x(fin, datos)",
        "cross join lateral unnest(string_to_array(x.datos, ',')) as kv",
        "join public.prod_semanas s on s.fecha_fin = x.fin",
        "join public.prod_operadoras o on o.folio_empleado = split_part(kv, ':', 1)",
        "on conflict (semana_id, operadora_id) do nothing;",
        "commit;",
        "",
    ]
    with open(OUT_SQL, "w", encoding="utf-8") as f:
        f.write("\n".join(sql))
    print("\n".join(lineas))
    print(f"\nfilas de valor a insertar: {sum(len(por_fecha[conservar[f]]) for f in conservar)} -> {OUT_SQL}")


if __name__ == "__main__":
    main()
